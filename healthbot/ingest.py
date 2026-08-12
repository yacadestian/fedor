"""Bulk history import: Yandex Disk folders + mail attachments → filtered,
deduplicated medical archive + one consolidated text document.

Pipeline per file: download → extract text (pdfplumber; scanned PDFs are
rendered page-by-page through the vision API) → LLM filter "medical?" →
medical files go to the bundle (raw + .txt), everything lands in manifest,
and the consolidated document is built from medical records sorted by date.

    .venv/bin/python import_history.py --yadisk-dir /Здоровье --mail \
        --out data/history_bundle
    .venv/bin/python import_history.py --from-bundle data/history_bundle --to-db
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
import tempfile
from dataclasses import dataclass, field, asdict
from pathlib import Path

log = logging.getLogger("health-bot")

MEDICAL_FILTER_PROMPT = """Определи, является ли документ медицинским. Верни СТРОГО JSON (без markdown):
{"medical": true/false,
 "kind": "lab_results|consultation|research|prescription|discharge|vaccination|dental|other_medical|not_medical",
 "title": "краткое название документа",
 "date": "YYYY-MM-DD или null (дата оказания услуги/сдачи анализа, не дата печати)",
 "summary": "одно предложение: что это и ключевая информация"}

Медицинские: анализы, заключения врачей, результаты исследований (МРТ/УЗИ/ЭЭГ/ЭКГ),
направления, назначения, выписки, справки о здоровье, прививки, стоматология.
НЕ медицинские: чеки магазинов, билеты, банковские выписки, договоры, счета,
государственные письма, рассылки — даже если от клиники, но без медицинского
содержания (просто оплата/реклама).

=== ТЕКСТ ДОКУМЕНТА ===
"""


@dataclass
class IngestRecord:
    file_hash: str
    filename: str
    source: str            # "yadisk" | "mail"
    source_ref: str        # disk path or "folder | subject | date"
    size_bytes: int
    medical: bool = False
    kind: str = ""
    title: str = ""
    doc_date: str = ""
    summary: str = ""
    text_chars: int = 0
    error: str = ""

    @property
    def sort_key(self) -> str:
        return self.doc_date or "9999"


def pdf_to_text(pdf_path: Path, max_pages_ocr: int = 10) -> str:
    """pdfplumber for text PDFs; scanned pages rendered through the vision API."""
    from .pdf_parser import extract_text
    text, _pages = extract_text(pdf_path)
    # Tiny embedded text layers (scanner junk) still need vision OCR
    if len(text.strip()) >= 80:
        return text
    ocr_text = scanned_pdf_to_text(pdf_path, max_pages_ocr)
    if len(ocr_text.strip()) > len(text.strip()):
        return ocr_text
    return text or ocr_text


def scanned_pdf_to_text(pdf_path: Path, max_pages: int = 10) -> str:
    """Render a scanned PDF page-by-page and OCR each page via vision API."""
    import time
    import fitz  # pymupdf
    from .ocr import _vision_api_read, VISION_API_MODEL_FAST

    texts: list[str] = []
    with fitz.open(pdf_path) as doc:
        for i, page in enumerate(doc):
            if i >= max_pages:
                log.info("Scanned PDF truncated at %d pages: %s", max_pages, pdf_path.name)
                break
            pix = page.get_pixmap(dpi=180)
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                pix.save(tmp.name)
                tmp_path = Path(tmp.name)
            try:
                # Prefer flash-lite: same quality for lab scans, much higher free quota
                page_text = _vision_api_read(tmp_path, VISION_API_MODEL_FAST)
            finally:
                tmp_path.unlink(missing_ok=True)
            if page_text:
                texts.append(page_text)
            time.sleep(2.0)  # stay under Gemini free-tier RPM
    return "\n\n".join(texts)


def classify_medical(text: str) -> dict:
    """LLM filter: is this a medical document? Strict JSON out."""
    from .llm import chat as llm_chat
    answer = llm_chat(MEDICAL_FILTER_PROMPT + text[:4000],
                      tier="fast", timeout=60, temperature=0.0)
    m = re.search(r"\{.*\}", answer, re.DOTALL)
    if not m:
        raise ValueError(f"Нет JSON в ответе фильтра: {answer[:200]}")
    return json.loads(m.group(0))


_JUNK_SUFFIXES = {".css", ".js", ".html", ".htm", ".svg", ".ico", ".map", ".woff",
                  ".woff2", ".ttf", ".zip", ".gz", ".mp3", ".mp4", ".wav"}
_UI_NAME_RE = re.compile(
    r"(^|[_-])(media_|section_|icon_|sprite_|thumb)|"
    r"_thumb\.|@2x\.|@3x\.",
    re.IGNORECASE,
)


def _is_junk_asset(filename: str, data: bytes) -> str | None:
    """Skip Telegram/web UI crumbs without spending vision API quota."""
    suffix = Path(filename).suffix.lower()
    name = Path(filename).name
    if suffix in _JUNK_SUFFIXES:
        return f"skip junk: {suffix}"
    if suffix in (".jpg", ".jpeg", ".png", ".webp", ".gif"):
        if len(data) < 20_000:  # <20 KB — icons, not lab photos
            return "skip tiny image"
        if _UI_NAME_RE.search(name):
            return "skip ui asset"
    return None


def _docx_to_text(data: bytes) -> str:
    import io
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def process_file(filename: str, data: bytes, source: str, source_ref: str,
                 work_dir: Path, seen: set[str]) -> IngestRecord:
    """Full pipeline for one file. Never raises — errors land in the record."""
    digest = hashlib.sha256(data).hexdigest()
    rec = IngestRecord(file_hash=digest, filename=filename, source=source,
                       source_ref=source_ref, size_bytes=len(data))
    if digest in seen:
        rec.error = "duplicate"
        return rec
    seen.add(digest)

    junk = _is_junk_asset(filename, data)
    if junk:
        rec.error = junk
        rec.kind = "not_medical"
        return rec

    try:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            work_dir.mkdir(parents=True, exist_ok=True)
            tmp_pdf = work_dir / f"{digest[:16]}.pdf"
            tmp_pdf.write_bytes(data)
            text = pdf_to_text(tmp_pdf)
        elif suffix in (".jpg", ".jpeg", ".png", ".webp"):
            from .ocr import _vision_api_read, VISION_API_MODEL_FAST
            tmp_img = work_dir / f"{digest[:16]}{suffix}"
            tmp_img.write_bytes(data)
            # API-only OCR via Gemini flash-lite (no local Tesseract)
            text = _vision_api_read(tmp_img, VISION_API_MODEL_FAST)
        elif suffix == ".docx":
            text = _docx_to_text(data)
        elif suffix == ".txt":
            text = data.decode("utf-8", errors="replace")
        else:
            rec.error = f"unsupported type: {suffix}"
            return rec
    except Exception as exc:
        rec.error = f"extract: {exc}"
        log.warning("Extract failed %s: %s", filename, exc)
        return rec

    rec.text_chars = len(text)
    if len(text.strip()) < 30:
        rec.error = "no text"
        return rec

    try:
        verdict = classify_medical(text)
        rec.medical = bool(verdict.get("medical"))
        # LLM may return JSON null — coalesce so logging/slicing never crash
        rec.kind = verdict.get("kind") or ""
        rec.title = verdict.get("title") or ""
        rec.doc_date = verdict.get("date") or ""
        rec.summary = verdict.get("summary") or ""
    except Exception as exc:
        rec.error = f"classify: {exc}"
        log.warning("Classify failed %s: %s", filename, exc)
        return rec

    # Always persist recognized text; raw originals only for medical docs
    (work_dir / "text").mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\wа-яА-ЯёЁ.-]+", "_", filename)[:120]
    txt_name = f"{digest[:8]}_{Path(safe).stem}.txt"
    (work_dir / "text" / txt_name).write_text(text, encoding="utf-8")
    _save_text_cache(work_dir, digest, text)
    if rec.medical:
        (work_dir / "raw").mkdir(parents=True, exist_ok=True)
        (work_dir / "raw" / f"{digest[:8]}_{safe}").write_bytes(data)
    return rec


def _save_text_cache(work_dir: Path, digest: str, text: str) -> None:
    (work_dir / ".textcache").mkdir(exist_ok=True)
    (work_dir / ".textcache" / f"{digest}.txt").write_text(text, encoding="utf-8")


def _load_text_cache(work_dir: Path, digest: str) -> str:
    p = work_dir / ".textcache" / f"{digest}.txt"
    return p.read_text(encoding="utf-8") if p.exists() else ""


_SOURCE_LABELS = {"yadisk": "Яндекс.Диск", "mail": "Почта"}


def build_consolidated(records: list[IngestRecord], work_dir: Path) -> str:
    """One text document with all medical records, sorted by document date."""
    medical = sorted([r for r in records if r.medical], key=lambda r: r.sort_key)
    skipped = [r for r in records if not r.medical and not r.error]
    lines = [
        "# Сводная история анализов и обследований",
        "",
        f"Документов медицинских: {len(medical)} | отфильтровано немедицинских: {len(skipped)}",
        "",
    ]
    current_year = None
    for r in medical:
        year = (r.doc_date or "")[:4] or "без даты"
        if year != current_year:
            current_year = year
            lines.append(f"\n{'='*60}\n## {year}\n{'='*60}\n")
        src = _SOURCE_LABELS.get(r.source, r.source)
        lines.append(f"\n--- {r.doc_date or 'дата неизвестна'} | {r.title or r.filename}")
        lines.append(f"Источник: {src} | {r.source_ref}")
        if r.summary:
            lines.append(f"Суть: {r.summary}")
        lines.append("")
        text = _load_text_cache(work_dir, r.file_hash)
        lines.append(text[:15000] if text else "(текст не найден)")
        lines.append("")
    return "\n".join(lines)


def save_manifest(records: list[IngestRecord], work_dir: Path) -> Path:
    work_dir.mkdir(parents=True, exist_ok=True)
    path = work_dir / "manifest.json"
    path.write_text(json.dumps([asdict(r) for r in records],
                               ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_manifest(work_dir: Path) -> list[IngestRecord]:
    data = json.loads((work_dir / "manifest.json").read_text(encoding="utf-8"))
    return [IngestRecord(**r) for r in data]


def import_bundle_to_db(work_dir: Path, owner_id: str) -> tuple[int, int]:
    """Push a collected bundle into ClickHouse: documents + lab biomarkers."""
    from .db import insert_document, insert_lab_results
    from .extractor import extract_biomarkers, validate_results
    from datetime import date

    records = [r for r in load_manifest(work_dir) if r.medical and not r.error]
    n_docs = n_biomarkers = 0
    for r in records:
        text = _load_text_cache(work_dir, r.file_hash)
        if not text:
            continue
        try:
            doc_date = date.fromisoformat(r.doc_date) if r.doc_date else date.today()
        except ValueError:
            doc_date = date.today()
        kind = r.kind or "other"
        if kind == "lab_results":
            try:
                extracted = extract_biomarkers(text)
                rows, _ = validate_results(extracted)
                for row in rows:
                    row["source_file"] = r.filename
                    row["raw_text"] = text[:10000]
                n_biomarkers += insert_lab_results(rows, owner_id)
                continue
            except Exception as exc:
                log.warning("Biomarker extraction failed for %s, saving as document: %s",
                            r.filename, exc)
        insert_document(collected_at=doc_date, doc_type=kind, title=r.title or r.filename,
                        source_file=f"{r.source}:{r.filename}", full_text=text,
                        owner_id=owner_id)
        n_docs += 1
    return n_docs, n_biomarkers
